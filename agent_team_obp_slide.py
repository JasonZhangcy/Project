# -*- coding: utf-8 -*-
"""生成《Agent Team 竞争力构建规划》OBP 单页 PPT。
左侧:分层规划架构图;右侧:竞争力特性规划表。
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 配色 ----------
HW_RED = RGBColor(0xC7, 0x00, 0x0B)        # 华为红(强调色)
DARK = RGBColor(0x1F, 0x2A, 0x3A)          # 深灰蓝(正文/标题)
GRAY = RGBColor(0x5A, 0x64, 0x72)          # 次要文字
L1 = RGBColor(0x2C, 0x5F, 0x9E)            # 体验层
L2 = RGBColor(0xC7, 0x00, 0x0B)            # 调度层(差异化主打,用红色)
L3 = RGBColor(0x1F, 0x7A, 0x6D)            # 编排层
L4 = RGBColor(0x4A, 0x55, 0x68)            # 底座
ENTRY = RGBColor(0x8A, 0x2B, 0xE2) if False else RGBColor(0x6B, 0x4F, 0xA0)  # 入口
SUB_FILL = RGBColor(0xFF, 0xFF, 0xFF)
TBL_HDR = RGBColor(0x1F, 0x2A, 0x3A)
TBL_ROW_A = RGBColor(0xF5, 0xF7, 0xFA)
TBL_ROW_B = RGBColor(0xEA, 0xEE, 0xF4)
FONT = "微软雅黑"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank


def set_font(run, size, color=DARK, bold=False, name=FONT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = name
    # 东亚字体需单独设置
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)


def add_text(shape, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    """lines: list of (text, size, color, bold)"""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(4)
    tf.margin_top = tf.margin_bottom = Pt(2)
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text
        set_font(r, size, color, bold)
    return shape


def box(x, y, w, h, fill, line=None, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    sp = slide.shapes.add_shape(shape_type, x, y, w, h)
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


# ================= 标题区 =================
title = slide.shapes.add_textbox(Inches(0.35), Inches(0.14), Inches(12.6), Inches(0.55))
add_text(title, [("Agent Team 竞争力构建规划:从单 Agent 到企业级数字团队", 26, DARK, True)],
         align=PP_ALIGN.LEFT)

subtitle = slide.shapes.add_textbox(Inches(0.38), Inches(0.66), Inches(12.6), Inches(0.34))
add_text(subtitle, [("统一入口 · 自动路由 · 自主编排 · 专家资产化 —— 三层递进能力 + 企业级治理底座,构筑差异化竞争力", 12.5, GRAY, False)],
         align=PP_ALIGN.LEFT)

# 标题下红色装饰线
bar = box(Inches(0.38), Inches(0.60), Inches(1.6), Pt(3), HW_RED,
          shape_type=MSO_SHAPE.RECTANGLE)

TOP = 1.12          # 内容区起始
LEFT_X = 0.35
LEFT_W = 5.95

# ================= 左侧:规划构建架构图 =================
sec1 = slide.shapes.add_textbox(Inches(LEFT_X), Inches(TOP - 0.04), Inches(LEFT_W), Inches(0.3))
add_text(sec1, [("规划构建架构", 15, HW_RED, True)], align=PP_ALIGN.LEFT)

ARCH_TOP = TOP + 0.32
LABEL_W = 0.92      # 左侧层名标签宽
GAP = 0.07
BODY_X = LEFT_X + LABEL_W + 0.08
BODY_W = LEFT_W - LABEL_W - 0.08

def layer(y, h, label, color, title_text, items, item_rows=1):
    # 层标签(纵向)
    lab = box(Inches(LEFT_X), Inches(y), Inches(LABEL_W), Inches(h), color)
    add_text(lab, [(label, 11.5, RGBColor(0xFF, 0xFF, 0xFF), True)])
    # 层主体
    body = box(Inches(BODY_X), Inches(y), Inches(BODY_W), Inches(h),
               RGBColor(0xF2, 0xF5, 0xF9), line=color)
    # 层标题
    t = slide.shapes.add_textbox(Inches(BODY_X + 0.08), Inches(y + 0.02),
                                 Inches(BODY_W - 0.16), Inches(0.26))
    add_text(t, [(title_text, 11, color, True)], align=PP_ALIGN.LEFT)
    # 子项小方块
    n = len(items)
    cols = (n + item_rows - 1) // item_rows
    iw = (BODY_W - 0.16 - (cols - 1) * 0.06) / cols
    ih = (h - 0.34 - (item_rows - 1) * 0.05) / item_rows
    for idx, it in enumerate(items):
        r, c = divmod(idx, cols)
        ix = BODY_X + 0.08 + c * (iw + 0.06)
        iy = y + 0.30 + r * (ih + 0.05)
        sp = box(Inches(ix), Inches(iy), Inches(iw), Inches(ih), SUB_FILL, line=color, radius=0.16)
        add_text(sp, [(it, 9.5, DARK, False)])


y = ARCH_TOP
# 统一任务入口
entry = box(Inches(BODY_X), Inches(y), Inches(BODY_W), Inches(0.44), ENTRY)
add_text(entry, [("统一任务入口:自然语言描述任务,免选单 Agent / Agent Team", 11.5, RGBColor(0xFF, 0xFF, 0xFF), True)])
lab0 = box(Inches(LEFT_X), Inches(y), Inches(LABEL_W), Inches(0.44), ENTRY)
add_text(lab0, [("入口", 11.5, RGBColor(0xFF, 0xFF, 0xFF), True)])
y += 0.44 + GAP

# 体验层
layer(y, 1.02, "体验层", L1, "专家 / 专家团(降低使用门槛,资产化运营)",
      ["预置领域\n专家库", "场景化专家\n团模板", "企业专家\n工厂", "专家/技能\n市场"])
y += 1.02 + GAP

# 调度层
layer(y, 1.02, "调度层", L2, "复杂度自适应路由(差异化主打:业界尚无显式产品化)",
      ["任务复杂度\n评估器", "三档执行\n模式决策", "运行中动态\n升降级", "成本-质量\n策略路由"])
y += 1.02 + GAP

# 编排层
layer(y, 1.02, "编排层", L3, "自主编排与固化(确定性、可复用)",
      ["递归Planner\n任务拆解", "Teammate\n动态合成", "团队蓝图\n固化", "例行化\n定时/事件触发"])
y += 1.02 + GAP

# 底座
layer(y, 1.02, "底座", L4, "企业级治理控制面(政企商用落地前提)",
      ["权限与\n审计", "成本配额\n管控", "效果度量\nROI 看板", "共享记忆\n上下文空间"])
y += 1.02

# 架构原则注脚
note = slide.shapes.add_textbox(Inches(LEFT_X), Inches(y + 0.05), Inches(LEFT_W), Inches(0.5))
add_text(note, [("架构原则:写操作单线程 + 多 Agent 贡献智能(map-reduce-and-manage),规避并行写冲突与成本失控", 9.5, GRAY, False)],
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

# ================= 右侧:竞争力特性规划表 =================
RIGHT_X = 6.55
RIGHT_W = 6.45
sec2 = slide.shapes.add_textbox(Inches(RIGHT_X), Inches(TOP - 0.04), Inches(RIGHT_W), Inches(0.3))
add_text(sec2, [("竞争力特性规划", 15, HW_RED, True)], align=PP_ALIGN.LEFT)

rows_data = [
    ("专家/专家团资产体系",
     "预置研发领域专家库与场景化专家团模板,一句话唤起整团;企业专家工厂支持规范/资产蒸馏定制;专家市场实现共享、版本化与审核治理"),
    ("任务自主编排(自动组队)",
     "递归 Planner 意图理解与任务拆解,按子任务自动合成 Teammate(角色/模型/工具/上下文);共享任务列表 + 依赖图调度,并行推进"),
    ("团队蓝图固化",
     "编队结构、角色分工、执行流程一键沉淀为参数化、版本化蓝图资产,支持重放与回归评测,解决\u201c每次生成不一致\u201d;支持用户自组团队并固化"),
    ("复杂度自适应路由",
     "统一入口免选择,复杂度评估驱动三档路由(单 Agent / 子 Agent / 团队),决策可解释、可覆盖;运行中动态升降级,业界首个显式产品化"),
    ("成本-质量协同治理",
     "团队模式 Token 成本随规模线性增长,路由决策纳入成本预算约束;基于任务级反馈在线学习,持续优化成本-质量帕累托前沿"),
    ("企业级治理与度量底座",
     "统一控制面纳管全部 Agent 资产:权限、审计、配额;路由决策全程留痕可审计;Agent 效果与 ROI 度量看板;跨任务共享记忆空间"),
]

tbl_top = TOP + 0.32
tbl_h = 5.62
gfx = slide.shapes.add_table(len(rows_data) + 1, 2,
                             Inches(RIGHT_X), Inches(tbl_top), Inches(RIGHT_W), Inches(tbl_h))
table = gfx.table
table.columns[0].width = Inches(1.72)
table.columns[1].width = Inches(RIGHT_W - 1.72)
table.rows[0].height = Inches(0.40)
for i in range(1, len(rows_data) + 1):
    table.rows[i].height = Inches((tbl_h - 0.40) / len(rows_data))

# 表头
hdr = ["关键竞争力", "竞争力描述"]
for c in range(2):
    cell = table.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = TBL_HDR
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = cell.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = hdr[c]
    set_font(r, 13, RGBColor(0xFF, 0xFF, 0xFF), True)

for i, (k, v) in enumerate(rows_data, start=1):
    fill = TBL_ROW_A if i % 2 == 1 else TBL_ROW_B
    # 第一列
    c0 = table.cell(i, 0)
    c0.fill.solid(); c0.fill.fore_color.rgb = fill
    c0.vertical_anchor = MSO_ANCHOR.MIDDLE
    c0.margin_left = Pt(5); c0.margin_right = Pt(3)
    p = c0.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = k
    accent = HW_RED if "路由" in k else DARK
    set_font(r, 11, accent, True)
    # 第二列
    c1 = table.cell(i, 1)
    c1.fill.solid(); c1.fill.fore_color.rgb = fill
    c1.vertical_anchor = MSO_ANCHOR.MIDDLE
    c1.margin_left = Pt(6); c1.margin_right = Pt(5)
    c1.text_frame.word_wrap = True
    p = c1.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = v
    set_font(r, 10, DARK, False)

# 右侧注脚:对标信息
foot = slide.shapes.add_textbox(Inches(RIGHT_X), Inches(tbl_top + tbl_h + 0.05), Inches(RIGHT_W), Inches(0.5))
add_text(foot, [("对标:Claude Code Agent Teams / Cursor 递归编排 / Devin Managed Devins / Grok Bot / WorkBuddy 专家团 / GitHub Agent HQ", 9, GRAY, False)],
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

prs.save("/workspace/Agent_Team_OBP规划.pptx")
print("saved")
