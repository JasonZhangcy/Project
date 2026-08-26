#!/usr/bin/env python3
"""Render the 16:9 HTML slide and wrap it into a one-page PPTX."""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
HTML = ROOT / "slides" / "madao-multi-end-obp.html"
OUT_DIR = ROOT / "output"
PNG_2X = OUT_DIR / "slide_2x.png"
PNG_1X = OUT_DIR / "slide_1920x1080.png"
PPTX = OUT_DIR / "码道多端协同能力规划.pptx"

NOTES = """【口播】五端不是五套产品，是同一套码道 Agent 的五种工作表面。

左：场景 → 五端定位 → 端云协同五步（起草、上云、长程、驾驭、回灌），中枢是统一任务总线 + Runtime + CodeArts Repo。
右：先做中枢指标（可见 100%、热迁移 ≥95%、回灌 ≥90%、Token -30%），再分端落地。

关键区分：
- IDE = 人机结对驾驶舱；CLI = Agent 进程与流水线接口，不是 IDE 精简版。
- Space = 泛研发/泛办公工作台，开发任务出口给 Web/IDE，不做成第二 IDE。
- Web = 云端沙箱 + Repo 为中心（对标 Origin，长在 CodeArts）。
- Mobile = 驾驭与审批面（微信/元服务/App），不是迷你 IDE。

协同差异：友商已能上云；码道要把对话上下文 + dirty tree 一起带走，再一键回灌。
"""


def render_png() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    cmd = [
        "google-chrome",
        "--headless",
        "--disable-gpu",
        "--no-sandbox",
        "--disable-dev-shm-usage",
        "--hide-scrollbars",
        "--window-size=1920,1080",
        "--force-device-scale-factor=2",
        "--virtual-time-budget=5000",
        f"--screenshot={PNG_2X}",
        HTML.resolve().as_uri(),
    ]
    try:
        subprocess.run(cmd, check=True, timeout=60)
    except subprocess.TimeoutExpired:
        if not PNG_2X.exists():
            raise
    im = Image.open(PNG_2X)
    if im.size != (3840, 2160):
        print(f"warn: unexpected size {im.size}", file=sys.stderr)
    im.resize((1920, 1080), Image.Resampling.LANCZOS).save(PNG_1X, optimize=True)


def add_notes(slide, text: str) -> None:
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = text.strip()
    for p in tf.paragraphs:
        p.font.size = Pt(12)
        p.font.name = "Microsoft YaHei"
        p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        str(PNG_2X),
        Emu(0),
        Emu(0),
        width=prs.slide_width,
        height=prs.slide_height,
    )
    add_notes(slide, NOTES)
    prs.save(PPTX)


def main() -> None:
    render_png()
    build_pptx()
    print(f"wrote {PNG_2X}")
    print(f"wrote {PNG_1X}")
    print(f"wrote {PPTX}")


if __name__ == "__main__":
    main()
