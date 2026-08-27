#!/usr/bin/env python3
"""Generate the CodeArts Agent Team OBP one-pager (PPTX + PNG preview)."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
OUT_PPTX = OUT_DIR / "码道AgentTeam竞争力构建规划.pptx"
OUT_PNG = OUT_DIR / "码道AgentTeam竞争力构建规划.png"

# Slide geometry (16:9)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

FONT_NAME = "微软雅黑"
CJK_FONT = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"
FALLBACK_FONT = "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf"

# Palette — Huawei Cloud / OBP
RED = "C7000B"
NAVY = "14233D"
BLUE = "1B5FA8"
TEAL = "1F7A6B"
GOLD = "9A7209"
PURPLE = "5A4E8C"
ENGINE = "2C3E50"
GRAY = "5B6573"
LINE = "D7DDE5"
BG = "F4F6F8"
WHITE = "FFFFFF"
SOFT_BLUE = "E8F1FB"
SOFT_TEAL = "E7F4F1"
SOFT_GOLD = "F8F1DE"
SOFT_PURPLE = "EFEBF8"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def emu_px(inches: float, dpi: int = 144) -> int:
    return int(round(inches * dpi))


def set_run_font(run, size_pt, bold=False, color_hex=WHITE, name=FONT_NAME):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = name
    run.font.color.rgb = rgb(color_hex)
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        node = rPr.find(qn(f"a:{tag}"))
        if node is None:
            node = etree.SubElement(rPr, qn(f"a:{tag}"))
        node.set("typeface", name)


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=12,
    bold=False,
    color=NAVY,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.MIDDLE,
    font_name=FONT_NAME,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.clear()
    run = p.add_run()
    run.text = text
    set_run_font(run, size, bold, color, font_name)
    return box


def set_shape_text(
    shape,
    lines,
    size=11,
    bold=False,
    color=WHITE,
    align=PP_ALIGN.CENTER,
    anchor="ctr",
    size_sub=10,
    color_sub=None,
):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf._txBody.bodyPr.set("anchor", anchor)
    if isinstance(lines, str):
        lines = [(lines, size, bold, color)]
    for i, item in enumerate(lines):
        if isinstance(item, str):
            txt, sz, bd, col = item, (size if i == 0 else size_sub), (bold if i == 0 else False), (color if i == 0 else (color_sub or color))
        else:
            txt, sz, bd, col = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        p.clear()
        run = p.add_run()
        run.text = txt
        set_run_font(run, sz, bd, col)


def style_shape(shape, fill, line=None, line_w=1.0, radius=0.12):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_w)
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass


def add_round(slide, left, top, width, height, fill, line=None, radius=0.12):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    style_shape(shape, fill, line, radius=radius)
    return shape


def add_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    style_shape(shape, fill, line, radius=0)
    return shape


def add_down_chevron(slide, cx, top, color=BLUE):
    w, h = Inches(0.16), Inches(0.09)
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, cx - w / 2, top, w, h)
    style_shape(shape, color, radius=0)
    return shape


def set_cell(cell, text, size, bold, color, fill, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.vertical_anchor = valign
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.10)
    tf.margin_right = Inches(0.10)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(4 if i == 0 else 2)
        p.clear()
        run = p.add_run()
        run.text = line
        set_run_font(run, size if i == 0 else max(size - 0.5, 9), bold if i == 0 else False, color)
        if i > 0:
            set_run_font(run, 10.5, False, color)


def build_pptx():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.08), RED)
    add_rect(slide, 0, Inches(7.26), SLIDE_W, Inches(0.24), NAVY)

    # Title
    add_textbox(
        slide, Inches(0.32), Inches(0.16), Inches(10.6), Inches(0.38),
        "Agent Team 竞争力构建规划", size=24, bold=True, color=NAVY,
    )
    add_textbox(
        slide, Inches(0.32), Inches(0.52), Inches(12.6), Inches(0.28),
        "以统一入口驱动专家资产化、编排标准化与执行拓扑自主化，将协作引擎升级为可规模交付的研发作业体系",
        size=12, color=GRAY,
    )

    # Section titles
    add_rect(slide, Inches(0.32), Inches(0.90), Inches(0.08), Inches(0.22), RED)
    add_textbox(slide, Inches(0.46), Inches(0.86), Inches(5.6), Inches(0.30),
                "规划构建架构", size=14, bold=True, color=NAVY)
    add_rect(slide, Inches(6.72), Inches(0.90), Inches(0.08), Inches(0.22), RED)
    add_textbox(slide, Inches(6.86), Inches(0.86), Inches(5.9), Inches(0.30),
                "竞争力特性规划", size=14, bold=True, color=NAVY)

    # ---- Left architecture ----
    ax, aw = Inches(0.32), Inches(6.16)
    tag_w = Inches(0.70)
    box_x = ax + tag_w + Inches(0.08)
    box_w = aw - tag_w - Inches(0.08)
    y = Inches(1.18)
    gap = Inches(0.10)

    def layer_tag(top, h, text, fill):
        s = add_round(slide, ax, top, tag_w, h, fill, radius=0.16)
        set_shape_text(s, text, size=10, bold=True, color=WHITE)

    # 体验层
    h = Inches(0.72)
    layer_tag(y, h, "体验层", RED)
    s = add_round(slide, box_x, y, box_w, h, RED, radius=0.10)
    set_shape_text(s, [("统一任务入口", 14, True, WHITE), ("用户无需预选单 Agent / Agent Team", 11, False, "F8D0D4")])
    y = y + h + gap
    add_down_chevron(slide, box_x + box_w / 2, y - gap, RED)

    # 调度层
    h = Inches(1.02)
    layer_tag(y, h, "调度层", BLUE)
    s = add_round(slide, box_x, y, box_w, Inches(0.50), BLUE, radius=0.10)
    set_shape_text(s, [("执行拓扑路由器  Topology Router", 13, True, WHITE)])
    pill_y = y + Inches(0.58)
    pills = [("Solo", "EAF2FB", BLUE), ("Subagent", "EAF2FB", BLUE), ("专家团", "EAF2FB", BLUE), ("Best-of-N", "EAF2FB", BLUE)]
    pill_gap = Inches(0.08)
    pw = (box_w - pill_gap * 3) / 4
    for i, (name, bg, fg) in enumerate(pills):
        p = add_round(slide, box_x + i * (pw + pill_gap), pill_y, pw, Inches(0.36), bg, line=BLUE, radius=0.20)
        set_shape_text(p, name, size=11, bold=True, color=fg)
    y = y + h + gap
    add_down_chevron(slide, box_x + box_w / 2, y - gap, BLUE)

    # 编排层 — dual path
    h = Inches(1.22)
    layer_tag(y, h, "编排层", TEAL)
    half = (box_w - Inches(0.10)) / 2
    left = add_round(slide, box_x, y, half, h, SOFT_TEAL, line=TEAL, radius=0.10)
    set_shape_text(
        left,
        [("命中 Playbook", 12, True, TEAL), ("固化专家团执行", 14, True, NAVY), ("确定性拓扑 · 质量门禁 · 版本一致", 10, False, GRAY)],
        align=PP_ALIGN.CENTER,
        color=NAVY,
    )
    right = add_round(slide, box_x + half + Inches(0.10), y, half, h, SOFT_GOLD, line=GOLD, radius=0.10)
    set_shape_text(
        right,
        [("未命中 → 动态编排", 12, True, GOLD), ("Intent-to-Team", 14, True, NAVY), ("拆解 Teammate / DAG，成功后固化沉淀", 10, False, GRAY)],
        align=PP_ALIGN.CENTER,
        color=NAVY,
    )
    y = y + h + gap
    add_down_chevron(slide, box_x + box_w / 2, y - gap, TEAL)

    # 资产层
    h = Inches(0.84)
    layer_tag(y, h, "资产层", PURPLE)
    s = add_round(slide, box_x, y, box_w, h, PURPLE, radius=0.10)
    set_shape_text(s, [("专家中心 / 企业专家库 / 个人专家", 13, True, WHITE), ("人设 + 方法论 + 工具链 + 验收标准", 11, False, "E4DDF3")])
    y = y + h + gap
    add_down_chevron(slide, box_x + box_w / 2, y - gap, PURPLE)

    # 引擎层
    h = Inches(0.83)
    layer_tag(y, h, "引擎层", ENGINE)
    s = add_round(slide, box_x, y, box_w, h, ENGINE, radius=0.10)
    set_shape_text(s, [("Agent Team 协作运行时", 13, True, WHITE), ("Leader 智能编排 · Teammate 自主执行 · 双向通信 · DAG 可视", 10, False, "D5DDE6")])
    y = y + h + gap
    add_down_chevron(slide, box_x + box_w / 2, y - gap, ENGINE)

    # 治理层
    h = Inches(0.83)
    layer_tag(y, h, "治理层", NAVY)
    s = add_round(slide, box_x, y, box_w, h, NAVY, radius=0.10)
    set_shape_text(s, [("质量闭环与企业管控", 13, True, WHITE), ("强制检视席 · CodeArts 流水线门禁 · 审计留痕 · Token/并发配额", 10, False, "C9D3E0")])

    # ---- Right table ----
    rows, cols = 4, 2
    table_shape = slide.shapes.add_table(rows, cols, Inches(6.72), Inches(1.18), Inches(6.28), Inches(5.96))
    table = table_shape.table
    table.columns[0].width = Inches(1.72)
    table.columns[1].width = Inches(4.56)

    # Set row heights
    table.rows[0].height = Inches(0.42)
    for i in range(1, 4):
        table.rows[i].height = Inches(1.847)

    headers = ("关键竞争力", "竞争力描述")
    for j, htxt in enumerate(headers):
        set_cell(table.cell(0, j), htxt, 12, True, WHITE, RED, align=PP_ALIGN.CENTER)

    data = [
        (
            "研发专家/\n专家团资产化",
            "将 Agent Team 产品化为可召唤的研发岗位资产，以「人设 + 方法论 + 工具链 + 验收标准」沉淀专家能力。\n"
            "• 构建官方预置 / 企业专家库 / 个人专家三级货架，一键召唤、隐藏编排细节\n"
            "• 预置 Issue 修复、三方库升级、安全检视、需求到交付等高频专家团\n"
            "• 对标 WorkBuddy 消费体验，差异化绑定软件工程闭环与企业资产治理",
        ),
        (
            "自主编排与\nPlaybook 固化",
            "基于任务描述自动识别意图，约束化拆解 Teammate 并生成可执行 DAG，避免角色随机漂移。\n"
            "• 升级 /save-team 为 Playbook：固化角色、流程骨架、工具绑定、门禁与版本\n"
            "• 支持用户自组团（可视化 + 自然语言）并沉淀为企业可复用模板\n"
            "• 同类任务二次执行拓扑一致，解决“每次生成团队不同”的稳态问题",
        ),
        (
            "执行拓扑\n自动路由",
            "提供统一任务入口，按范围、可并行度、角色异构、风险与预算自动选择执行拓扑。\n"
            "• 四级路由：Solo / Subagent / 专家团 / Best-of-N，默认 Solo、满足条件再升级\n"
            "• 运行中支持升降级；路由结果可解释、可人工覆盖，并纳入企业配额策略\n"
            "• 友商仅有模型路由或 Subagent 委派，拓扑路由尚无完整产品化，具备窗口期领先性",
        ),
    ]
    fills = (SOFT_BLUE, "F7F8FA", SOFT_BLUE)
    title_colors = (BLUE, TEAL, RED)
    for i, ((title, desc), fill) in enumerate(zip(data, fills), start=1):
        set_cell(table.cell(i, 0), title, 12, True, title_colors[i - 1], fill, align=PP_ALIGN.CENTER)
        # description with mixed paragraphs
        cell = table.cell(i, 1)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
        cell.vertical_anchor = MSO_ANCHOR.TOP
        tf = cell.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.10)
        tf.margin_top = Inches(0.10)
        tf.margin_bottom = Inches(0.08)
        paras = desc.split("\n")
        for k, line in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(0)
            p.space_after = Pt(5 if k == 0 else 2)
            p.clear()
            run = p.add_run()
            run.text = line
            set_run_font(run, 11 if k == 0 else 10.5, k == 0, NAVY if k == 0 else "3D4654")

    # Footer
    add_textbox(
        slide, Inches(0.32), Inches(7.28), Inches(9.2), Inches(0.20),
        "华为云码道 CodeArts  |  OBP 规划  |  对标 WorkBuddy / Grok Bot / Claude Code / Copilot CLI",
        size=10, color=WHITE, anchor=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide, Inches(10.2), Inches(7.28), Inches(2.8), Inches(0.20),
        "内部公开",
        size=10, color=WHITE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)
    return OUT_PPTX


def load_font(size: int, index: int = 0) -> ImageFont.FreeTypeFont:
    for path in (CJK_FONT, FALLBACK_FONT):
        try:
            return ImageFont.truetype(path, size, index=index)
        except Exception:
            continue
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_w: int) -> list[str]:
    if not text:
        return [""]
    lines, cur = [], ""
    for ch in text:
        trial = cur + ch
        if draw.textlength(trial, font=font) <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = ch
    if cur:
        lines.append(cur)
    return lines or [text]


def rounded(draw, xy, r, fill, outline=None, width=1):
    draw.rounded_rectangle(xy, radius=r, fill=fill, outline=outline, width=width)


def hex_to_rgb(h: str) -> tuple[int, int, int]:
    return tuple(int(h[i : i + 2], 16) for i in (0, 2, 4))


def center_text(draw, text, cx, cy, font, fill):
    bbox = draw.textbbox((0, 0), text, font=font)
    w, h = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.text((cx - w / 2 - bbox[0], cy - h / 2 - bbox[1]), text, font=font, fill=fill)


def build_png():
    W, H = 1920, 1080
    img = Image.new("RGB", (W, H), hex_to_rgb(WHITE))
    d = ImageDraw.Draw(img)

    font_title = load_font(36)
    font_sub = load_font(18)
    font_sec = load_font(22)
    font_box_title = load_font(20)
    font_box_sub = load_font(14)
    font_tag = load_font(14)
    font_pill = load_font(15)
    font_th = load_font(18)
    font_comp = load_font(18)
    font_desc = load_font(16)
    font_bullet = load_font(15)
    font_foot = load_font(13)

    d.rectangle([0, 0, W, 10], fill=hex_to_rgb(RED))
    d.rectangle([0, 1046, W, H], fill=hex_to_rgb(NAVY))

    d.text((46, 24), "Agent Team 竞争力构建规划", font=font_title, fill=hex_to_rgb(NAVY))
    d.text(
        (46, 72),
        "以统一入口驱动专家资产化、编排标准化与执行拓扑自主化，将协作引擎升级为可规模交付的研发作业体系",
        font=font_sub,
        fill=hex_to_rgb(GRAY),
    )

    # section titles
    d.rectangle([46, 126, 56, 154], fill=hex_to_rgb(RED))
    d.text((66, 124), "规划构建架构", font=font_sec, fill=hex_to_rgb(NAVY))
    d.rectangle([968, 126, 978, 154], fill=hex_to_rgb(RED))
    d.text((988, 124), "竞争力特性规划", font=font_sec, fill=hex_to_rgb(NAVY))

    # architecture geometry
    tag_x, tag_w = 46, 92
    box_x, box_w = 148, 770
    y = 170
    layer_gap = 14

    def chevron(cx, y0, color):
        d.polygon([(cx - 8, y0), (cx + 8, y0), (cx, y0 + 10)], fill=hex_to_rgb(color))

    def tag_box(top, h, text, fill):
        rounded(d, [tag_x, top, tag_x + tag_w, top + h], 10, hex_to_rgb(fill))
        center_text(d, text, tag_x + tag_w / 2, top + h / 2, font_tag, hex_to_rgb(WHITE))

    # 体验层
    h = 104
    tag_box(y, h, "体验层", RED)
    rounded(d, [box_x, y, box_x + box_w, y + h], 12, hex_to_rgb(RED))
    center_text(d, "统一任务入口", box_x + box_w / 2, y + h / 2 - 14, font_box_title, hex_to_rgb(WHITE))
    center_text(d, "用户无需预选单 Agent / Agent Team", box_x + box_w / 2, y + h / 2 + 16, font_box_sub, (248, 208, 212))
    y = y + h + layer_gap
    chevron(box_x + box_w / 2, y - 12, RED)

    # 调度层
    h = 147
    tag_box(y, h, "调度层", BLUE)
    rounded(d, [box_x, y, box_x + box_w, y + 72], 12, hex_to_rgb(BLUE))
    center_text(d, "执行拓扑路由器  Topology Router", box_x + box_w / 2, y + 36, font_box_title, hex_to_rgb(WHITE))
    pills = ["Solo", "Subagent", "专家团", "Best-of-N"]
    gap, pw = 10, (box_w - 30) / 4
    py = y + 86
    for i, name in enumerate(pills):
        px = box_x + i * (pw + gap)
        rounded(d, [px, py, px + pw, py + 46], 18, hex_to_rgb(SOFT_BLUE), outline=hex_to_rgb(BLUE), width=1)
        center_text(d, name, px + pw / 2, py + 23, font_pill, hex_to_rgb(BLUE))
    y = y + h + layer_gap
    chevron(box_x + box_w / 2, y - 12, BLUE)

    # 编排层
    h = 176
    tag_box(y, h, "编排层", TEAL)
    half = (box_w - 14) / 2
    rounded(d, [box_x, y, box_x + half, y + h], 12, hex_to_rgb(SOFT_TEAL), outline=hex_to_rgb(TEAL), width=2)
    center_text(d, "命中 Playbook", box_x + half / 2, y + 42, font_pill, hex_to_rgb(TEAL))
    center_text(d, "固化专家团执行", box_x + half / 2, y + 82, font_box_title, hex_to_rgb(NAVY))
    center_text(d, "确定性拓扑 · 质量门禁 · 版本一致", box_x + half / 2, y + 122, font_box_sub, hex_to_rgb(GRAY))
    rx = box_x + half + 14
    rounded(d, [rx, y, rx + half, y + h], 12, hex_to_rgb(SOFT_GOLD), outline=hex_to_rgb(GOLD), width=2)
    center_text(d, "未命中 → 动态编排", rx + half / 2, y + 42, font_pill, hex_to_rgb(GOLD))
    center_text(d, "Intent-to-Team", rx + half / 2, y + 82, font_box_title, hex_to_rgb(NAVY))
    center_text(d, "拆解 Teammate / DAG，成功后固化沉淀", rx + half / 2, y + 122, font_box_sub, hex_to_rgb(GRAY))
    y = y + h + layer_gap
    chevron(box_x + box_w / 2, y - 12, TEAL)

    # 资产层
    h = 121
    tag_box(y, h, "资产层", PURPLE)
    rounded(d, [box_x, y, box_x + box_w, y + h], 12, hex_to_rgb(PURPLE))
    center_text(d, "专家中心 / 企业专家库 / 个人专家", box_x + box_w / 2, y + h / 2 - 14, font_box_title, hex_to_rgb(WHITE))
    center_text(d, "人设 + 方法论 + 工具链 + 验收标准", box_x + box_w / 2, y + h / 2 + 16, font_box_sub, (228, 221, 243))
    y = y + h + layer_gap
    chevron(box_x + box_w / 2, y - 12, PURPLE)

    # 引擎层
    h = 120
    tag_box(y, h, "引擎层", ENGINE)
    rounded(d, [box_x, y, box_x + box_w, y + h], 12, hex_to_rgb(ENGINE))
    center_text(d, "Agent Team 协作运行时", box_x + box_w / 2, y + h / 2 - 14, font_box_title, hex_to_rgb(WHITE))
    center_text(d, "Leader 智能编排 · Teammate 自主执行 · 双向通信 · DAG 可视", box_x + box_w / 2, y + h / 2 + 16, font_box_sub, (213, 221, 230))
    y = y + h + layer_gap
    chevron(box_x + box_w / 2, y - 12, ENGINE)

    # 治理层
    h = 120
    tag_box(y, h, "治理层", NAVY)
    rounded(d, [box_x, y, box_x + box_w, y + h], 12, hex_to_rgb(NAVY))
    center_text(d, "质量闭环与企业管控", box_x + box_w / 2, y + h / 2 - 14, font_box_title, hex_to_rgb(WHITE))
    center_text(d, "强制检视席 · CodeArts 流水线门禁 · 审计留痕 · Token/并发配额", box_x + box_w / 2, y + h / 2 + 16, font_box_sub, (201, 211, 224))

    # ---- table ----
    tx, ty, tw, th = 968, 170, 906, 858
    col0 = 248
    header_h = 52
    row_h = (th - header_h) / 3

    rounded(d, [tx, ty, tx + tw, ty + th], 8, hex_to_rgb(WHITE), outline=hex_to_rgb(LINE), width=1)
    d.rectangle([tx, ty, tx + tw, ty + header_h], fill=hex_to_rgb(RED))
    center_text(d, "关键竞争力", tx + col0 / 2, ty + header_h / 2, font_th, hex_to_rgb(WHITE))
    center_text(d, "竞争力描述", tx + col0 + (tw - col0) / 2, ty + header_h / 2, font_th, hex_to_rgb(WHITE))
    d.line([tx + col0, ty, tx + col0, ty + th], fill=hex_to_rgb(LINE), width=1)

    table_rows = [
        (
            ["研发专家/", "专家团资产化"],
            BLUE,
            SOFT_BLUE,
            "将 Agent Team 产品化为可召唤的研发岗位资产，以「人设 + 方法论 + 工具链 + 验收标准」沉淀专家能力。",
            [
                "构建官方预置 / 企业专家库 / 个人专家三级货架，一键召唤、隐藏编排细节",
                "预置 Issue 修复、三方库升级、安全检视、需求到交付等高频专家团",
                "对标 WorkBuddy 消费体验，差异化绑定软件工程闭环与企业资产治理",
            ],
        ),
        (
            ["自主编排与", "Playbook 固化"],
            TEAL,
            "F7F8FA",
            "基于任务描述自动识别意图，约束化拆解 Teammate 并生成可执行 DAG，避免角色随机漂移。",
            [
                "升级 /save-team 为 Playbook：固化角色、流程骨架、工具绑定、门禁与版本",
                "支持用户自组团（可视化 + 自然语言）并沉淀为企业可复用模板",
                "同类任务二次执行拓扑一致，解决“每次生成团队不同”的稳态问题",
            ],
        ),
        (
            ["执行拓扑", "自动路由"],
            RED,
            SOFT_BLUE,
            "提供统一任务入口，按范围、可并行度、角色异构、风险与预算自动选择执行拓扑。",
            [
                "四级路由：Solo / Subagent / 专家团 / Best-of-N，默认 Solo、满足条件再升级",
                "运行中支持升降级；路由结果可解释、可人工覆盖，并纳入企业配额策略",
                "友商仅有模型路由或 Subagent 委派，拓扑路由尚无完整产品化，具备窗口期领先性",
            ],
        ),
    ]

    for i, (title_lines, tcolor, fill, lead, bullets) in enumerate(table_rows):
        ry = ty + header_h + i * row_h
        d.rectangle([tx + 1, ry, tx + tw - 1, ry + row_h - (0 if i < 2 else 1)], fill=hex_to_rgb(fill))
        d.line([tx, ry, tx + tw, ry], fill=hex_to_rgb(LINE), width=1)
        d.line([tx + col0, ry, tx + col0, ry + row_h], fill=hex_to_rgb(LINE), width=1)
        # title
        total_h = len(title_lines) * 26
        ty0 = ry + (row_h - total_h) / 2
        for k, line in enumerate(title_lines):
            center_text(d, line, tx + col0 / 2, ty0 + 13 + k * 26, font_comp, hex_to_rgb(tcolor))
        # desc
        dx, dw = tx + col0 + 16, tw - col0 - 32
        desc_y = ry + 16
        for line in wrap_text(d, lead, font_desc, dw):
            d.text((dx, desc_y), line, font=font_desc, fill=hex_to_rgb(NAVY))
            desc_y += 22
        desc_y += 6
        for b in bullets:
            wrapped = wrap_text(d, "•  " + b, font_bullet, dw)
            for wline in wrapped:
                d.text((dx, desc_y), wline, font=font_bullet, fill=hex_to_rgb("3D4654"))
                desc_y += 20
            desc_y += 2

    d.text(
        (46, 1056),
        "华为云码道 CodeArts  |  OBP 规划  |  对标 WorkBuddy / Grok Bot / Claude Code / Copilot CLI",
        font=font_foot,
        fill=hex_to_rgb(WHITE),
    )
    fw = d.textlength("内部公开", font=font_foot)
    d.text((W - 46 - fw, 1056), "内部公开", font=font_foot, fill=hex_to_rgb(WHITE))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    img.save(OUT_PNG, "PNG", optimize=True)
    return OUT_PNG


if __name__ == "__main__":
    pptx_path = build_pptx()
    png_path = build_png()
    print(f"PPTX: {pptx_path}")
    print(f"PNG:  {png_path}")
