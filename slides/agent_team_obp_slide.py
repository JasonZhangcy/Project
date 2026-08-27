"""生成华为云码道 CodeArts Agent Team 规划 OBP 单页胶片（16:9，原生可编辑形状 + 表格）。

左侧：规划构建架构图（任务执行主链路 + 形态分支 + 团队内部结构 + 资产/底座注入与反馈回环）
右侧：关键竞争力特性表

用法：python3 slides/agent_team_obp_slide.py [输出路径.pptx]
"""

import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

LATIN_FONT = "Arial"
EA_FONT = "微软雅黑"

INK = RGBColor(0x14, 0x21, 0x3D)
MUTED = RGBColor(0x5A, 0x66, 0x78)
HAIRLINE = RGBColor(0xD9, 0xE0, 0xEA)
CARD_BG = RGBColor(0xF6, 0xF8, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BAND = RGBColor(0xF6, 0xF8, 0xFB)

BLUE = RGBColor(0x2E, 0x6B, 0xE6)
BLUE_BG = RGBColor(0xEE, 0xF3, 0xFE)
ACCENT = RGBColor(0xC7, 0x00, 0x0B)
ACCENT_BG = RGBColor(0xFD, 0xF1, 0xF1)
PURPLE = RGBColor(0x7A, 0x3E, 0xBF)
PURPLE_BG = RGBColor(0xF4, 0xEF, 0xFB)
TEAL = RGBColor(0x0F, 0x9B, 0x8E)
TEAL_BG = RGBColor(0xED, 0xF7, 0xF6)
GRAY = RGBColor(0x4A, 0x55, 0x68)
GRAY_BG = RGBColor(0xF2, 0xF4, 0xF7)
FLOW = RGBColor(0x6B, 0x7A, 0x90)
DOTTED = RGBColor(0x94, 0xA3, 0xB8)

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.42
CONTENT_TOP = 1.72
CONTENT_BOTTOM = 6.80
LEFT_X, LEFT_W = MARGIN, 5.66
RIGHT_X = LEFT_X + LEFT_W + 0.33
RIGHT_W = SLIDE_W - MARGIN - RIGHT_X

TITLE = "Agent Team 规划：从多智能体引擎 → 智能体组织生产力平台"
SUBTITLE = (
    "三层跃迁：能力可召唤（专家化）· 编排可复现（资产化）· 形态可自主（路由化）"
    "；底座：团队记忆复利 + 可信治理与成本可控"
)
EYEBROW = "华为云码道 CodeArts · OBP 规划"

# ---------- 左侧架构图布局 ----------
RAIL_W = 1.02
RAIL_GAP = 0.18
MAIN_X = LEFT_X + RAIL_W + RAIL_GAP
MAIN_W = LEFT_W - 2 * (RAIL_W + RAIL_GAP)
CX = MAIN_X + MAIN_W / 2
LRAIL_X = LEFT_X
RRAIL_X = LEFT_X + LEFT_W - RAIL_W
RAIL_Y, RAIL_H = 2.32, 3.06

ASSET_CHIPS = ["Role Profile", "Blueprint 三态", "场景专家团包", "市场·分级治理", "Codebase 注入"]
OPS_CHIPS = ["记忆·轨迹复盘", "沙箱·单线程", "预算·配额·熔断", "审计·权限治理", "可观测大盘"]
FORMS = [("单 Agent", False), ("单 Agent + 评审", False), ("Agent Team", True)]
GATES = ["CodeCheck", "CloudTest", "流水线门禁"]
TEAMMATES = ["Teammate·分析", "Teammate·开发", "Teammate·测试"]

TABLE_HEADER = ("关键竞争力", "竞争力描述")
TABLE_ROWS = [
    (
        "专家化封装\n零门槛召唤",
        "Role Profile 契约化 + 场景专家团预置包，以\u201c专家 / 专家团\u201d形态交付多智能体能力；"
        "对标 WorkBuddy 三层心智，叠加码道 Skills 与 Codebase",
        False,
    ),
    (
        "团队资产固化\n与版本治理",
        "Blueprint 三态（Draft→Blueprint→Certified）+ 引用版本锁定 + 参数化槽位，"
        "破解\u201c每次拆解都不一样\u201d，上线拆解一致性指标；对标 Qoder 与 Copilot 的版本锁定实践",
        False,
    ),
    (
        "自主编排\n与自主组队",
        "Leader 自动识别任务、动态拆解并按需建组；拓扑 DSL 与可视化编辑支持用户自建团队并一键固化；"
        "Trace→Blueprint 自动萃取，把成功轨迹沉淀为组织级资产",
        False,
    ),
    (
        "编排形态自主路由\n（战略高地）",
        "统一入口下由 Harness 判定单体 / 单体+评审 / 小团队 / 全量团队；"
        "业界仅到子 Agent 自动委派（L2），形态级路由（L3）尚无产品化闭环，学术已验证降本至多 52%",
        True,
    ),
    (
        "升配降配\n与成本护栏",
        "默认单体起步，触发条件下 checkpoint 移交团队续跑而非重启，收敛后自动降配；"
        "任务与租户级预算、递归深度与并发上限、异常熔断，支撑大客户成本可承诺",
        False,
    ),
    (
        "团队记忆\n与经验复利",
        "Teammate 持续化上下文 + 执行轨迹复盘 + 失败模式库反哺路由与拆解；"
        "Codebase 与企业知识按角色定向注入，形成\u201c越用越强\u201d，区别于竞品一次性构造的团队",
        False,
    ),
    (
        "企业级可信\n与门禁闭环",
        "单线程写入 + worktree 沙箱隔离；产出必过 CodeCheck / CloudTest / 流水线门禁；"
        "最小权限、全链路审计、四级作用域治理与私有化合规——纯 IDE 类竞品的结构性短板",
        False,
    ),
    (
        "场景化\n差异化资产",
        "鸿蒙适配迁移专家团 + 存量代码 Codebase 深度理解 + 华为研发体系 Skills 沉淀，"
        "构筑竞品不可迁移的场景护城河",
        False,
    ),
]

FOOTER_LINES = [
    "度量口径：拆解一致性 · 蓝图复用率 · 路由准确率 · Token/任务 · 人工干预次数 · 一次通过率 · 并行加速比",
    "数据锚点：多智能体 ≈ 15× Chat Token（Anthropic）· 多智能体轨迹失败率 41%~87%（NeurIPS'25）· "
    "编排形态路由降本至多 52%（MasRouter, ACL'25）",
]


# ---------------- 基础绘制helper ----------------
def style_run(run, size, bold=False, color=INK, space=None):
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    font.name = LATIN_FONT
    rPr = run._r.get_or_add_rPr()
    if space is not None:
        rPr.set("spc", str(int(space * 100)))
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", EA_FONT)


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=0.95, space=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    para = frame.paragraphs[0]
    para.alignment = align
    para.line_spacing = line
    run = para.add_run()
    run.text = text
    style_run(run, size, bold, color, space)
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    box = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    box.shadow.inherit = False
    if fill is None:
        box.fill.background()
    else:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
        box.line.width = Pt(0.75)
    if radius is not None and box.adjustments:
        box.adjustments[0] = radius
    box.text_frame.word_wrap = True
    return box


def block(slide, x, y, w, h, specs, fill, line=None, radius=0.16,
          shape=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER, line_spacing=1.05):
    """带居中多行文字的图形块。specs = [(文本, 字号, 加粗, 颜色), ...]"""
    box = rect(slide, x, y, w, h, fill, line, shape=shape, radius=radius)
    frame = box.text_frame
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for idx, (text, size, bold, color) in enumerate(specs):
        para = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        para.alignment = align
        para.line_spacing = line_spacing
        run = para.add_run()
        run.text = text
        style_run(run, size, bold, color)
    return box


def down_arrow(slide, cx, y, h=0.14, w=0.15, color=FLOW):
    return rect(slide, cx - w / 2, y, w, h, color, shape=MSO_SHAPE.DOWN_ARROW)


def right_arrow(slide, x, cy, w=0.16, h=0.11, color=FLOW):
    return rect(slide, x, cy - h / 2, w, h, color, shape=MSO_SHAPE.RIGHT_ARROW)


def line_seg(slide, x1, y1, x2, y2, color=FLOW, weight=0.9):
    x, y = min(x1, x2), min(y1, y2)
    w, h = max(abs(x2 - x1), weight / 72), max(abs(y2 - y1), weight / 72)
    return rect(slide, x, y, w, h, color)


def dotted_arrow(slide, x1, y1, x2, y2, color=DOTTED):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(1.1)
    ln = conn.line._get_or_add_ln()
    dash = ln.makeelement(qn("a:prstDash"), {"val": "sysDash"})
    ln.append(dash)
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "sm", "len": "sm"})
    ln.append(tail)
    return conn


def set_cell(cell, text, size, bold=False, color=INK, fill=WHITE,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, line=1.06):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(0.075)
    cell.margin_right = Inches(0.075)
    cell.margin_top = Inches(0.045)
    cell.margin_bottom = Inches(0.045)
    cell.vertical_anchor = anchor
    frame = cell.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.alignment = align
    para.line_spacing = line
    run = para.add_run()
    run.text = text
    style_run(run, size, bold, color)


# ---------------- 左侧架构图 ----------------
def draw_rail(slide, x, header, chips, color, bg):
    rect(slide, x, RAIL_Y, RAIL_W, RAIL_H, bg, color,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    textbox(slide, x + 0.05, RAIL_Y + 0.09, RAIL_W - 0.10, 0.34,
            header, 8, True, color, align=PP_ALIGN.CENTER, line=1.1)
    chip_h, gap = 0.40, 0.13
    top = RAIL_Y + 0.48
    for idx, chip in enumerate(chips):
        block(slide, x + 0.07, top + idx * (chip_h + gap), RAIL_W - 0.14, chip_h,
              [(chip, 7, False, INK)], WHITE, HAIRLINE, radius=0.14)


def draw_architecture(slide):
    # 侧栏：能力资产 / 运维底座
    draw_rail(slide, LRAIL_X, "专家 / 专家团\n资产库", ASSET_CHIPS, TEAL, TEAL_BG)
    draw_rail(slide, RRAIL_X, "Agent Team Ops\n底座", OPS_CHIPS, GRAY, GRAY_BG)

    # ① 统一入口
    block(slide, MAIN_X + 0.35, 1.72, MAIN_W - 0.70, 0.42,
          [("① 统一入口 One Entry", 9.5, True, WHITE),
           ("自然语言任务 / 专家召唤 · IDE · CLI · 流水线", 7, False, WHITE)],
          BLUE, radius=0.36)
    down_arrow(slide, CX, 2.16)

    # ② 复杂度画像
    block(slide, MAIN_X, 2.32, MAIN_W, 0.44,
          [("② 任务复杂度画像", 9.5, True, ACCENT),
           ("任务信号 · 仓库耦合信号 · 可并行度 · 历史表现", 7, False, MUTED)],
          ACCENT_BG, ACCENT, radius=0.14)
    down_arrow(slide, CX, 2.78)

    # ③ 策略引擎（决策核心）
    block(slide, MAIN_X, 2.94, MAIN_W, 0.50,
          [("③ 编排策略引擎  ·  形态自主路由", 10, True, WHITE),
           ("形态 · 规模 · 模型档位 · 预算    ｜    L3 业界空白", 7, False, WHITE)],
          ACCENT, radius=0.10, shape=MSO_SHAPE.HEXAGON)

    # 分支：三种执行形态
    form_w, form_gap = (MAIN_W - 0.20) / 3, 0.10
    centers = [MAIN_X + form_w / 2 + i * (form_w + form_gap) for i in range(3)]
    line_seg(slide, CX, 3.44, CX, 3.50, FLOW)
    line_seg(slide, centers[0], 3.50, centers[2], 3.50, FLOW)
    for idx, (label, is_team) in enumerate(FORMS):
        line_seg(slide, centers[idx], 3.50, centers[idx], 3.56, FLOW)
        down_arrow(slide, centers[idx], 3.50, h=0.10, w=0.12)
        block(slide, MAIN_X + idx * (form_w + form_gap), 3.60, form_w, 0.38,
              [(label, 8, True, PURPLE if is_team else INK)],
              PURPLE_BG if is_team else WHITE, PURPLE if is_team else HAIRLINE,
              radius=0.18)

    # 升配 / 降配 双向关系
    rect(slide, centers[0], 4.04, centers[2] - centers[0], 0.16, PURPLE_BG, PURPLE,
         shape=MSO_SHAPE.LEFT_RIGHT_ARROW)
    textbox(slide, MAIN_X, 4.22, MAIN_W, 0.16,
            "升配 / 降配 · checkpoint 续跑（不重启）", 7, False, PURPLE,
            align=PP_ALIGN.CENTER)

    # ④ Agent Team 执行结构
    rect(slide, MAIN_X, 4.42, MAIN_W, 1.00, PURPLE_BG, PURPLE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    textbox(slide, MAIN_X + 0.10, 4.47, MAIN_W - 0.20, 0.18,
            "④ Agent Team 执行", 8, True, PURPLE)
    block(slide, CX - 0.55, 4.66, 1.10, 0.22,
          [("Team Leader", 7.5, True, WHITE)], PURPLE, radius=0.30)
    mate_w = 0.90
    mate_gap = (MAIN_W - 0.24 - 3 * mate_w) / 2
    mate_centers = [MAIN_X + 0.12 + mate_w / 2 + i * (mate_w + mate_gap) for i in range(3)]
    line_seg(slide, CX, 4.88, CX, 4.92, PURPLE)
    line_seg(slide, mate_centers[0], 4.92, mate_centers[2], 4.92, PURPLE)
    for idx, mate in enumerate(TEAMMATES):
        line_seg(slide, mate_centers[idx], 4.92, mate_centers[idx], 4.96, PURPLE)
        block(slide, mate_centers[idx] - mate_w / 2, 4.96, mate_w, 0.24,
              [(mate, 7, False, INK)], WHITE, PURPLE, radius=0.26)
    textbox(slide, MAIN_X + 0.10, 5.24, MAIN_W - 0.20, 0.14,
            "共享任务池 · 双向通信 · 写入单线程串行", 6.5, False, MUTED,
            align=PP_ALIGN.CENTER)
    down_arrow(slide, CX, 5.44)

    # ⑤ 质量门禁 → 交付
    gate_w = (MAIN_W - 0.32) / 3
    for idx, gate in enumerate(GATES):
        gx = MAIN_X + idx * (gate_w + 0.16)
        block(slide, gx, 5.58, gate_w, 0.32, [(gate, 7.5, True, GRAY)],
              GRAY_BG, GRAY, radius=0.20)
        if idx < 2:
            right_arrow(slide, gx + gate_w + 0.02, 5.74, w=0.12, color=GRAY)
    down_arrow(slide, CX, 5.92)
    block(slide, MAIN_X + 0.30, 6.08, MAIN_W - 0.60, 0.28,
          [("⑤ 交付件 · PR · 任务报告 · 度量", 8.5, True, INK)], WHITE, INK, radius=0.36)

    # 侧栏注入（虚线）
    for y in (3.19, 4.90):
        dotted_arrow(slide, LRAIL_X + RAIL_W, y, MAIN_X - 0.02, y, TEAL)
        dotted_arrow(slide, RRAIL_X, y, MAIN_X + MAIN_W + 0.02, y, GRAY)

    # 反馈回环（虚线）：执行轨迹 → 资产沉淀 / 路由学习
    feedback_y = 6.22
    line_x = LRAIL_X + RAIL_W / 2
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(MAIN_X + 0.30), Inches(feedback_y),
        Inches(line_x), Inches(feedback_y))
    conn.line.color.rgb = DOTTED
    conn.line.width = Pt(1.1)
    ln = conn.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:prstDash"), {"val": "sysDash"}))
    dotted_arrow(slide, line_x, feedback_y, line_x, RAIL_Y + RAIL_H + 0.04, TEAL)
    textbox(slide, LEFT_X, 6.40, LEFT_W, 0.16,
            "轨迹回流：Trace → Blueprint 自动萃取 · 失败模式库 · 路由策略学习",
            7, False, TEAL, align=PP_ALIGN.CENTER)
    textbox(slide, LEFT_X, 6.60, LEFT_W, 0.16,
            "实线＝任务执行主链路（①~⑤ P0 首版）　虚线＝资产 / 底座注入与反馈回环（P0 · P1 分批）",
            7, False, MUTED, align=PP_ALIGN.CENTER)


# ---------------- 右侧表格 ----------------
def draw_table(slide):
    header_h = 0.36
    frame = slide.shapes.add_table(
        len(TABLE_ROWS) + 1, 2, Inches(RIGHT_X), Inches(CONTENT_TOP),
        Inches(RIGHT_W), Inches(CONTENT_BOTTOM - CONTENT_TOP))
    table = frame.table
    tbl_pr = table._tbl.tblPr
    tbl_pr.set("firstRow", "0")
    tbl_pr.set("bandRow", "0")
    for style_id in tbl_pr.findall(qn("a:tableStyleId")):
        tbl_pr.remove(style_id)
    style_el = tbl_pr.makeelement(qn("a:tableStyleId"), {})
    style_el.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    tbl_pr.append(style_el)

    row_h = (CONTENT_BOTTOM - CONTENT_TOP - header_h) / len(TABLE_ROWS)
    table.columns[0].width = Inches(1.58)
    table.columns[1].width = Inches(RIGHT_W - 1.58)
    table.rows[0].height = Inches(header_h)
    for row in list(table.rows)[1:]:
        row.height = Inches(row_h)

    set_cell(table.cell(0, 0), TABLE_HEADER[0], 10.5, True, WHITE, INK,
             align=PP_ALIGN.CENTER)
    set_cell(table.cell(0, 1), TABLE_HEADER[1], 10.5, True, WHITE, INK)
    for i, (key, desc, highlight) in enumerate(TABLE_ROWS, start=1):
        if highlight:
            fill, key_color = ACCENT_BG, ACCENT
        else:
            fill, key_color = (WHITE if i % 2 else BAND), INK
        set_cell(table.cell(i, 0), key, 9.5, True, key_color, fill,
                 align=PP_ALIGN.CENTER)
        set_cell(table.cell(i, 1), desc, 8.5, False,
                 INK if highlight else MUTED, fill)


def build(path):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 标题区
    rect(slide, MARGIN, 0.30, 0.075, 0.52, ACCENT)
    textbox(slide, MARGIN + 0.20, 0.29, 10.4, 0.5, TITLE, 23, True, INK, space=0.4)
    textbox(slide, MARGIN + 0.20, 0.86, 11.6, 0.32, SUBTITLE, 11, False, MUTED)
    textbox(slide, SLIDE_W - MARGIN - 3.2, 0.34, 3.2, 0.26, EYEBROW, 10, False,
            MUTED, align=PP_ALIGN.RIGHT)
    rect(slide, MARGIN, 1.26, SLIDE_W - 2 * MARGIN, 0.012, HAIRLINE)

    # 栏目标题
    textbox(slide, LEFT_X, 1.36, LEFT_W, 0.3, "规划构建架构", 14, True, INK)
    rect(slide, LEFT_X, 1.655, 1.05, 0.032, ACCENT)
    textbox(slide, RIGHT_X, 1.36, RIGHT_W, 0.3, "竞争力特性规划", 14, True, INK)
    rect(slide, RIGHT_X, 1.655, 1.05, 0.032, ACCENT)

    draw_architecture(slide)
    draw_table(slide)

    # 页脚
    rect(slide, MARGIN, CONTENT_BOTTOM + 0.14, SLIDE_W - 2 * MARGIN, 0.012, HAIRLINE)
    footer = textbox(slide, MARGIN, CONTENT_BOTTOM + 0.23, SLIDE_W - 2 * MARGIN, 0.40,
                     FOOTER_LINES[0], 8.5, False, MUTED, line=1.25)
    second = footer.text_frame.add_paragraph()
    second.line_spacing = 1.25
    run = second.add_run()
    run.text = FOOTER_LINES[1]
    style_run(run, 8.5, False, MUTED)

    prs.save(path)
    return path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "slides/codearts-agent-team-obp.pptx"
    print(build(out))
