from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

FONT = "Microsoft YaHei"

RED = RGBColor(0xC0, 0x00, 0x00)
DARK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xF2, 0xA9, 0x00)
AMBER_DARK = RGBColor(0xB8, 0x78, 0x00)
AMBER_BG = RGBColor(0xFF, 0xF7, 0xE1)
AMBER_LINE = RGBColor(0xF5, 0xC9, 0x5C)
INTENT_BG = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x00, 0x8C, 0xD8)
BLUE_DARK = RGBColor(0x00, 0x5A, 0x9C)
BLUE_BG = RGBColor(0xE6, 0xF4, 0xFC)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def set_font(run, size, color=DARK, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def box(x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=0.75, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def text(x, y, w, h, paragraphs, anchor=MSO_ANCHOR.TOP, margin=0.04, align=PP_ALIGN.LEFT, spacing=1.12):
    """paragraphs: list of paragraphs; each paragraph is a list of (text, size, color, bold) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(1.5)
        for t, size, color, bold in runs:
            r = p.add_run()
            r.text = t
            set_font(r, size, color, bold)
    return tb


def est_lines(s, width_in, size):
    units = sum(0.55 if ord(c) < 128 else 1.0 for c in s)
    per_line = (width_in - 0.1) * 72 / size * 0.9
    lines, cur = 0, 0
    lines = max(1, -(-units // per_line))
    return int(lines)


def est_height(items, width_in, size, spacing=1.12):
    n = sum(est_lines(t, width_in, size) for t in items)
    return n * size * spacing * 1.2 / 72 + len(items) * 1.5 / 72 + 0.05


# ---------------- Title ----------------
box(0.35, 0.24, 0.08, 0.44, fill=RED)
text(0.5, 0.17, 12.6, 0.58, [[
    ("阿里云：Qoder、QoderWake、千问办公三线布局多人多Agent协作入口，驱动Agent负载上云", 20, RED, True)
]], anchor=MSO_ANCHOR.MIDDLE)
text(0.5, 0.72, 12.6, 0.3, [[
    ("核心判断：", 11, DARK, True),
    ("一套 Agent Harness 底座 + 三个协作入口（研发项目现场 / IM 数字员工 / 企业上下文），把 Agent 的“工作时长”转化为云上 CPU、存储、网络与 Token 消耗",
     11, GREY, False),
]], anchor=MSO_ANCHOR.MIDDLE)
line = slide.shapes.add_connector(1, Inches(0.35), Inches(1.06), Inches(12.98), Inches(1.06))
line.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
line.line.width = Pt(1)

# ---------------- Columns ----------------
COL_Y = 1.14
COL_H = 4.5
COL_W = 4.1
GAP = 0.17
COL_X = [0.35, 0.35 + COL_W + GAP, 0.35 + 2 * (COL_W + GAP)]
BODY = 9
LABEL = 9.5

columns = [
    {
        "name": "Qoder",
        "tag": "研发项目现场的人机协同",
        "meta": "9.23 云栖发布 · Teams / Enterprise 内测",
        "sections": [
            ("项目", [
                "以 Issue 拆解研发目标（目标 / 完成条件 / 优先级），6 态看板跟踪，关联代码仓",
                "Issue 一键发起讨论，背景自动带入，讨论资产沉淀在任务上",
            ]),
            ("讨论", [
                "成员与 Agent 同空间评审方案，@Agent 比较方案、补测试用例",
                "“关键消息”沉淀决策；Agent Team 由负责人 Agent 拆解、分派、汇总",
                "短板：仅能 @ 自有 Agent、需本机在线，结论需人工回写",
            ]),
        ],
        "intent": [
            "抢占“需求→代码”上游入口，决策留在哪、执行就留在哪",
            "协同仅企业版开放，推动个人订阅转向组织 Credits 池",
            "Agent / Skill 资产化，团队经验沉淀即迁移壁垒",
            "本机执行为过渡态，终局走向 Cloud Agents 云端托管",
        ],
    },
    {
        "name": "QoderWake",
        "tag": "IM 中 7×24 在岗的数字员工团队",
        "meta": "9 月发布 1.0 · 从“招一位”到“组建一支团队”",
        "sections": [
            ("多人多Agent协同", [
                "Waker 群组混编本地 / 远程 Waker，Leader 按 SOP 分派；WakerFlow 编排流程，定时 / 事件 / API 自动开工",
            ]),
            ("结合IM能力", [
                "@Waker 覆盖钉钉 / 飞书 / 企微等：理解历史消息、群任务、群记忆、话题整理，高危操作审批",
            ]),
            ("拥有长期记忆的数字员工", [
                "五层记忆：个人 / 项目 / 群 / 成员 / 单聊，自动沉淀 + 版本回滚",
                "Memory 与 Skill 自进化，成长时间线记录进化轨迹",
                "演进：补齐企业层、上云托管 → 个人-群-项目-企业四层记忆",
            ]),
        ],
        "intent": [
            "员工与工位分离：7×24、并发、合规需求把工位迁上云（无影 / ECS / 计算巢）",
            "常驻 Agent 持续消耗 CPU / 存储 / 网络：千人研发 × 2 Waker ≈ 4000 vCPU 常驻（示意）",
            "长期记忆 + 自进化形成切换成本，越用越难替换",
        ],
    },
    {
        "name": "千问办公",
        "tag": "企业上下文驱动的组织级协作",
        "meta": "9.22 云栖发布 · 与钉钉同一负责人",
        "sections": [
            ("企业上下文", [
                "汇聚 IM 与 SAP、Salesforce 数据，压缩 + 结构化 + 周期刷新，专属模型降 Token 成本",
                "公司级 / 部门级知识空间，Agent 按任务按需调用",
            ]),
            ("多人多Agent协同", [
                "协作空间关联群聊、文档、知识库，人-人、人-Agent、Agent-Agent 分工协作",
                "数字员工有部门、负责人、授权与生命周期；富立卡三员工接力",
                "Managed Agents 云端沙箱运行，安全中心全程审计",
            ]),
        ],
        "intent": [
            "上下文即护城河：成为企业数据面向 Agent 的统一出口",
            "复用钉钉组织身份与权限体系，锁定组织级入口",
            "上下文加工 + 托管运行，持续拉动数据处理与 Token 消耗",
        ],
    },
]

def sections_end(col):
    y = COL_Y + 0.64
    for label, items in col["sections"]:
        y += 0.24 + est_height(["▪ " + t for t in items], COL_W - 0.2, BODY) + 0.04
    return y


INTENT_Y = max(sections_end(c) for c in columns) + 0.02

for col, x in zip(columns, COL_X):
    box(x, COL_Y, COL_W, COL_H, fill=AMBER_BG, line=AMBER_LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
    hdr = box(x, COL_Y, COL_W, 0.56, fill=AMBER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    box(x, COL_Y + 0.4, COL_W, 0.16, fill=AMBER)
    text(x + 0.08, COL_Y + 0.02, COL_W - 0.16, 0.32, [[
        (col["name"], 14, WHITE, True), ("  |  " + col["tag"], 10.5, WHITE, True)
    ]], anchor=MSO_ANCHOR.MIDDLE)
    text(x + 0.08, COL_Y + 0.31, COL_W - 0.16, 0.24, [[(col["meta"], 8.5, RGBColor(0x5A, 0x3A, 0x00), False)]],
         anchor=MSO_ANCHOR.MIDDLE)

    y = COL_Y + 0.64
    inner_x = x + 0.1
    inner_w = COL_W - 0.2
    for label, items in col["sections"]:
        lw = 0.12 + sum(0.55 if ord(c) < 128 else 1.0 for c in label) * LABEL / 72 + 0.12
        tag = box(inner_x, y, lw, 0.22, fill=AMBER_DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
        text(inner_x, y, lw, 0.22, [[(label, LABEL, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER,
             margin=0.02)
        y += 0.24
        paras = [[("▪ ", BODY, AMBER_DARK, True), (t, BODY, DARK, False)] for t in items]
        h = est_height(["▪ " + t for t in items], inner_w, BODY)
        text(inner_x, y, inner_w, h, paras)
        y += h + 0.04

    intent_y = INTENT_Y
    intent_h = COL_Y + COL_H - 0.08 - intent_y
    box(inner_x, intent_y, inner_w, intent_h, fill=INTENT_BG, line=RGBColor(0xE8, 0x9C, 0x9C),
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    box(inner_x, intent_y + 0.06, 0.05, intent_h - 0.12, fill=RED)
    paras = [[("产品战略意图", LABEL + 0.5, RED, True)]]
    for i, t in enumerate(col["intent"], 1):
        paras.append([(f"{'①②③④'[i - 1]} ", BODY, RED, True), (t, BODY, DARK, False)])
    text(inner_x + 0.1, intent_y + 0.03, inner_w - 0.14, intent_h - 0.06, paras)

# ---------------- Insight bar ----------------
BAR_Y = 5.74
BAR_H = 1.47
box(0.35, BAR_Y, 12.63, BAR_H, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
text(0.5, BAR_Y + 0.04, 12.4, 0.34, [[
    ("观点（给码道 CodeArts 的启示）：", 12.5, WHITE, True),
    ("研发协作的竞争单元正从“个人 + Agent”转向“团队 + 数字员工”，码道应以项目为锚、以记忆为核、以云端工位变现", 11, WHITE, False),
]], anchor=MSO_ANCHOR.MIDDLE)

insights = [
    ("项目即上下文", "以 CodeArts 需求-代码-流水线-测试全链路对象为共享上下文，打通“讨论→拆单→Agent 执行→验收回写”闭环，补齐 Qoder 人工回写短板"),
    ("共享数字员工上云", "从“个人 Agent 进群”升级为项目级共享研发数字员工（需求 / 编码 / 检视 / 测试 / 值守），云端托管 7×24 在线，进驻 WeLink 等 IM"),
    ("四层记忆 + 持续进化", "构建个人-团队-项目-企业四层云端记忆，自动沉淀 + 人工确认 + 版本回滚；Skill 自进化经组织审核后统一下发，沉淀为企业研发资产"),
    ("数字员工拉动通算", "以“云端工位”承接数字员工常驻运行，按规格 × 时长计费，带动鲲鹏 CPU、存储、网络消耗，实现从“席位收入”到“席位 + 资源”双收入"),
]
card_w = (12.63 - 0.2 - 0.12 * 3) / 4
for i, (head, body) in enumerate(insights):
    cx = 0.45 + i * (card_w + 0.12)
    cy = BAR_Y + 0.42
    ch = BAR_H - 0.52
    box(cx, cy, card_w, ch, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    num = box(cx + 0.08, cy + 0.07, 0.28, 0.28, fill=BLUE_DARK, shape=MSO_SHAPE.OVAL)
    text(cx + 0.08, cy + 0.07, 0.28, 0.28, [[(str(i + 1), 10.5, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE,
         align=PP_ALIGN.CENTER, margin=0)
    text(cx + 0.4, cy + 0.05, card_w - 0.45, 0.32, [[(head, 11, BLUE_DARK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(cx + 0.08, cy + 0.37, card_w - 0.14, ch - 0.4, [[(body, BODY, DARK, False)]])

text(0.35, 7.23, 12.63, 0.2, [[(
    "资料来源：Qoder / QoderWake 官方文档与更新日志、Qoder Cloud Agents 1.0 发布文、2026 杭州云栖大会（吴泳铭、李飞飞主题演讲，千问办公发布会）公开报道；vCPU 规模为示意测算。",
    7, GREY, False)]], anchor=MSO_ANCHOR.MIDDLE)

prs.save("/workspace/insights/阿里云多人多Agent协作洞察.pptx")
print("saved")
