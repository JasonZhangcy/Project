# -*- coding: utf-8 -*-
"""生成《码道多端协同OBP规划》单页PPT(16:9,左右结构)。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
import copy

FONT = "Microsoft YaHei"

# 配色
HW_RED = RGBColor(0xC7, 0x00, 0x0B)      # 华为红(标题/强调)
DARK = RGBColor(0x1E, 0x29, 0x3B)        # 深色文字/中枢
GRAY = RGBColor(0x47, 0x55, 0x69)
LIGHT_BG = RGBColor(0xF4, 0xF6, 0xFA)
C_IDE = RGBColor(0x1F, 0x6F, 0xEB)       # 蓝
C_SPACE = RGBColor(0x7C, 0x3A, 0xED)     # 紫
C_WEB = RGBColor(0x0E, 0x80, 0x74)       # 青
C_MOBILE = RGBColor(0xD9, 0x77, 0x06)    # 橙
C_CLI = RGBColor(0x33, 0x41, 0x55)       # 深灰

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式


def set_text(tf, items, valign=MSO_ANCHOR.TOP):
    """items: list of (text, size, bold, color, align, space_after_pt)"""
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.02)
    for i, (text, size, bold, color, align, spa) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spa)
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn('a:ea'))
        if ea is None:
            ea = rPr.makeelement(qn('a:ea'), {})
            rPr.append(ea)
        ea.set('typeface', FONT)


def add_box(x, y, w, h, fill, line_color=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, shadow=False):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(1.0)
    sp.shadow.inherit = False
    return sp


def add_text(x, y, w, h, items, valign=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(tb.text_frame, items, valign)
    return tb


def add_arrow(x1, y1, x2, y2, color, width=1.5, dashed=False, both=False):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    ln = conn.line._get_or_add_ln()
    if dashed:
        dash = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
        ln.append(dash)
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    if both:
        head = ln.makeelement(qn('a:headEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
        ln.append(head)
    conn.shadow.inherit = False
    return conn


# ============ 标题栏 ============
add_box(0, 0, 13.333, 0.72, RGBColor(0xFF, 0xFF, 0xFF), shape=MSO_SHAPE.RECTANGLE)
add_box(0.25, 0.16, 0.09, 0.42, HW_RED, shape=MSO_SHAPE.RECTANGLE)
add_text(0.42, 0.08, 9.6, 0.6, [
    ("多端协同:一个大脑,五个界面 —— 任务随人走,算力在云端", 19, True, DARK, PP_ALIGN.LEFT, 0),
], MSO_ANCHOR.MIDDLE)
add_text(9.6, 0.08, 3.5, 0.6, [
    ("码道 OBP 规划 · 多端协同能力", 10, False, GRAY, PP_ALIGN.RIGHT, 0),
], MSO_ANCHOR.MIDDLE)
add_box(0.25, 0.74, 12.83, 0.012, RGBColor(0xD9, 0xDE, 0xE7), shape=MSO_SHAPE.RECTANGLE)

# ============ 左侧面板 ============
LX, LW = 0.25, 7.35

# ---- 1. 典型使用场景 ----
add_text(LX, 0.82, 3.0, 0.3, [("典型使用场景", 11, True, HW_RED, PP_ALIGN.LEFT, 0)])
scenarios = [
    ("下班不断线", "IDE任务一键上云,通勤中Mobile语音纠偏,到家手机审批合并PR"),
    ("会议室里派活", "PM在Space拆解Spec派发云端,开发者IDE同步可见、拉回精修"),
    ("流水线自愈", "CI失败,CLI无头模式自动修复开PR,Web总控可见、Mobile批准"),
    ("全民开发", "运营在Space用自然语言+技能做出应用,云端Agent完成部署"),
]
sc_w = (LW - 0.24) / 4
for i, (t, d) in enumerate(scenarios):
    x = LX + i * (sc_w + 0.08)
    add_box(x, 1.12, sc_w, 0.92, LIGHT_BG, RGBColor(0xD9, 0xDE, 0xE7), radius=0.12)
    add_text(x + 0.03, 1.16, sc_w - 0.06, 0.84, [
        ("场景%d %s" % (i + 1, t), 9, True, DARK, PP_ALIGN.LEFT, 2),
        (d, 7, False, GRAY, PP_ALIGN.LEFT, 0),
    ])

# ---- 2. 多端布局与协同架构 ----
add_text(LX, 2.14, 5.0, 0.3, [("多端布局与协同架构", 11, True, HW_RED, PP_ALIGN.LEFT, 0)])

BW, BH = 2.32, 1.42          # 端块尺寸
ROW1_Y, ROW2_Y = 2.48, 5.42  # 上下两行
HUB_X, HUB_Y, HUB_W, HUB_H = 2.62, 4.05, 2.6, 1.22

ends = [
    # (x, y, 颜色, 名称, 定位, 特性)
    (LX + 0.05, ROW1_Y, C_IDE, "码道 IDE", "专业开发者 · 深度创作主场",
     "Agent-first工作台 | 代码知识引擎\n多Agent并行 | Spec驱动+自主验证"),
    (LX + LW - BW - 0.05, ROW1_Y, C_SPACE, "码道 Space", "泛研发/泛办公 · 一站式AI工作台",
     "专家/专家团 | 项目空间 | 技能市场\nComputer/Browser Use | 深度研究/多模态"),
    (LX + 0.05, ROW2_Y, C_WEB, "码道 Web", "Repo为中心的云端Agent工场",
     "云端Agent并行 | 云端沙箱 | 长程执行\nAutomation/Merge Queue | 任务总控台"),
    (LX + (LW - BW) / 2, ROW2_Y, C_CLI, "码道 CLI", "终端原生的Agent引擎",
     "无头exec/脚本化 | CI/CD集成\nSDK可嵌入 | 企业级沙箱/审计"),
    (LX + LW - BW - 0.05, ROW2_Y, C_MOBILE, "码道 Mobile", "随身的Agent指挥端",
     "任务发起(语音) | 远程控制\n审批/PR合并 | 鸿蒙流转/小程序"),
]

# 中枢
hub = add_box(HUB_X, HUB_Y, HUB_W, HUB_H, DARK, radius=0.1)
add_text(HUB_X + 0.05, HUB_Y + 0.08, HUB_W - 0.1, HUB_H - 0.14, [
    ("码道云端任务中枢", 12, True, RGBColor(0xFF, 0xFF, 0xFF), PP_ALIGN.CENTER, 3),
    ("统一任务/会话模型 · 上下文与记忆", 7.5, False, RGBColor(0xC9, 0xD4, 0xE3), PP_ALIGN.CENTER, 1),
    ("云端沙箱 · Repo中心 · 知识/技能库", 7.5, False, RGBColor(0xC9, 0xD4, 0xE3), PP_ALIGN.CENTER, 0),
], MSO_ANCHOR.MIDDLE)

# 协同规则条(上行中间空档)
add_box(HUB_X, ROW1_Y + 0.06, HUB_W, 1.28, LIGHT_BG, RGBColor(0xD9, 0xDE, 0xE7), radius=0.1)
add_text(HUB_X + 0.04, ROW1_Y + 0.12, HUB_W - 0.08, 1.2, [
    ("三条协同规则", 8.5, True, HW_RED, PP_ALIGN.CENTER, 2),
    ("① 任务全局可见:任一端发起,全端同步", 7, False, DARK, PP_ALIGN.LEFT, 1),
    ("② 双向迁移:本地⇋云端一键接力", 7, False, DARK, PP_ALIGN.LEFT, 1),
    ("③ 人机分工:决策在人/执行在云/交互在端", 7, False, DARK, PP_ALIGN.LEFT, 0),
])

# 端块
for x, y, color, name, pos, feats in ends:
    add_box(x, y, BW, BH, RGBColor(0xFF, 0xFF, 0xFF), color, radius=0.09)
    add_box(x, y, BW, 0.30, color, radius=0.35)
    add_text(x + 0.02, y + 0.01, BW - 0.04, 0.28, [
        (name, 10, True, RGBColor(0xFF, 0xFF, 0xFF), PP_ALIGN.CENTER, 0)], MSO_ANCHOR.MIDDLE)
    lines = [(pos, 7.5, True, color, PP_ALIGN.CENTER, 2)]
    for seg in feats.split("\n"):
        lines.append((seg, 7, False, GRAY, PP_ALIGN.CENTER, 1))
    add_text(x + 0.03, y + 0.34, BW - 0.06, BH - 0.38, lines)

# 箭头:实线=发起/迁移(双向),虚线=同步/审批
ARROW = RGBColor(0x64, 0x74, 0x8B)
# IDE(上左) ↔ 中枢
add_arrow(1.85, ROW1_Y + BH, 2.95, HUB_Y, HW_RED, 1.8, both=True)
add_arrow(1.35, ROW1_Y + BH, 2.66, HUB_Y + 0.35, ARROW, 1.6, dashed=True)
# Space(上右) ↔ 中枢
add_arrow(6.0, ROW1_Y + BH, 4.9, HUB_Y, HW_RED, 1.8, both=True)
add_arrow(6.5, ROW1_Y + BH, 5.2, HUB_Y + 0.35, ARROW, 1.6, dashed=True)
# Web(下左) ↔ 中枢
add_arrow(1.85, ROW2_Y, 2.95, HUB_Y + HUB_H, HW_RED, 1.8, both=True)
add_arrow(1.35, ROW2_Y, 2.66, HUB_Y + HUB_H - 0.35, ARROW, 1.6, dashed=True)
# Mobile(下右) ↔ 中枢
add_arrow(6.0, ROW2_Y, 4.9, HUB_Y + HUB_H, HW_RED, 1.8, both=True)
add_arrow(6.5, ROW2_Y, 5.2, HUB_Y + HUB_H - 0.35, ARROW, 1.6, dashed=True)
# CLI(下中) ↔ 中枢
add_arrow(3.7, ROW2_Y, 3.7, HUB_Y + HUB_H, HW_RED, 1.8, both=True)
add_arrow(4.2, ROW2_Y, 4.2, HUB_Y + HUB_H, ARROW, 1.6, dashed=True)

# 图例 + 高亮协同路径
add_text(LX, 6.92, LW, 0.52, [
    ("—— 实线:任务发起/本地⇋云端迁移      ---- 虚线:同步观察/插话纠偏/审批", 7.5, False, GRAY, PP_ALIGN.LEFT, 2),
    ("路径示例① IDE本地会话→一键上云继续执行→Mobile通知审批    ② Web发起→Space同步可见→IDE拉回精修", 7.5, True, DARK, PP_ALIGN.LEFT, 0),
])

# ============ 右侧面板:竞争力规划特性 ============
RX, RW = 7.78, 5.3
add_box(RX - 0.08, 0.82, RW + 0.16, 6.55, LIGHT_BG, radius=0.03)
add_text(RX, 0.9, RW, 0.32, [("竞争力规划特性(落地项与关键指标)", 12.5, True, HW_RED, PP_ALIGN.LEFT, 0)])

# 主打:多端协同
y = 1.3
add_box(RX, y, RW, 0.98, DARK, radius=0.08)
add_text(RX + 0.06, y + 0.05, RW - 0.12, 0.9, [
    ("多端协同(主打)", 10, True, RGBColor(0xFF, 0xD7, 0x66), PP_ALIGN.LEFT, 2),
    ("· 任一端发起任务,全端100%可见可续,状态同步<3s", 8.5, False, RGBColor(0xFF, 0xFF, 0xFF), PP_ALIGN.LEFT, 1),
    ("· 本地⇋云端双向无损迁移<30s(上下文/变更/环境完整保留)· 开放会话协议接入三方客户端", 8.5, False, RGBColor(0xFF, 0xFF, 0xFF), PP_ALIGN.LEFT, 0),
])

groups = [
    (C_IDE, "码道 IDE", [
        "代码生成一次通过率≥80%,Spec+验证闭环下≥90%;同任务Token消耗降低≥50%",
        "代码知识引擎:仓库知识自动生成覆盖率≥95%,大仓任务准确率提升≥20%;单人并行≥5个Agent",
    ]),
    (C_SPACE, "码道 Space", [
        "专家团多Agent并行,复杂任务交付周期缩短≥60%;技能市场首年≥1000技能/≥100领域专家",
        "Computer/Browser Use任务成功率≥85%;DeepResearch报告一次可用率≥80%;多模态成果交付",
    ]),
    (C_WEB, "码道 Web", [
        "云端沙箱预构建启动<60s;长程任务自主执行≥24h、断点续跑、100%产出验证Artifact",
        "Repo中心:代码库实时镜像同步、事件驱动Automation、Agent感知Merge Queue、PR自动评审",
    ]),
    (C_MOBILE, "码道 Mobile", [
        "任务发起→云端执行<10s;审批推送触达<5s;远程控制本机IDE/CLI会话",
        "鸿蒙元服务卡片+跨设备流转(业界独有);微信小程序零安装查看/审批",
    ]),
    (C_CLI, "码道 CLI", [
        "无头模式+Agent SDK,CI/CD开箱即用;流水线失败自动修复率≥50%",
        "企业级:分级沙箱/命令审批策略/审计日志100%覆盖;Skills/MCP等扩展与全端互通",
    ]),
]
y = 2.42
GH = 0.94
for color, name, bullets in groups:
    add_box(RX, y, 0.06, GH - 0.1, color, shape=MSO_SHAPE.RECTANGLE)
    lines = [(name, 9.5, True, color, PP_ALIGN.LEFT, 2)]
    for b in bullets:
        lines.append(("· " + b, 8.5, False, DARK, PP_ALIGN.LEFT, 1))
    add_text(RX + 0.12, y - 0.04, RW - 0.14, GH, lines)
    y += GH + 0.045

prs.save("码道多端协同OBP规划.pptx")
print("saved")
